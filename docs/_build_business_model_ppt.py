# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt
from lxml import etree
from pathlib import Path

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0xF7, 0xF7, 0xF4)
INK = RGBColor(0x1C, 0x19, 0x17)
MUTED = RGBColor(0x57, 0x53, 0x4E)
BLUE = RGBColor(0x1A, 0x7F, 0xD4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE7, 0xE5, 0xE4)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
WARN = RGBColor(0xB4, 0x53, 0x09)
DARK = RGBColor(0x0C, 0x0A, 0x09)
FONT = "Microsoft YaHei"


def rgb_hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_run(run, text, size=14, color=INK, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)


def add_text(slide, l, t, w, h, text, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, color, bold)
    return box


def add_para(tf, text, size=14, color=INK, bold=False, space_after=6, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0]
    if p.runs:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size, color, bold)
    return p


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.adjustments[0] = 0.08
    fill_shape(s, color)
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    return s


def bar(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(s, color)
    return s


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bar(slide, Inches(0), Inches(0), Inches(0.12), H, BLUE)
    return slide


def footer(slide, page, total=12):
    add_text(slide, Inches(0.55), Inches(7.12), Inches(8), Inches(0.28), "光途Work · 商业模式  ·  内部讨论稿", 11, MUTED)
    add_text(slide, Inches(11.4), Inches(7.12), Inches(1.4), Inches(0.28), f"{page} / {total}", 11, MUTED, align=PP_ALIGN.RIGHT)


def heading(slide, title, subtitle=None):
    add_text(slide, Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.5), title, 26, INK, True)
    if subtitle:
        add_text(slide, Inches(0.55), Inches(0.82), Inches(12.2), Inches(0.36), subtitle, 13, MUTED)


def card_block(slide, l, t, w, h, title, body, accent=False):
    r = rect(slide, l, t, w, h, CARD, BLUE if accent else LINE)
    add_text(slide, l + Inches(0.22), t + Inches(0.16), w - Inches(0.4), Inches(0.36), title, 15, BLUE if accent else INK, True)
    box = slide.shapes.add_textbox(l + Inches(0.22), t + Inches(0.52), w - Inches(0.4), h - Inches(0.68))
    tf = box.text_frame
    tf.word_wrap = True
    lines = [ln for ln in body.split("\n") if ln]
    tf.paragraphs[0].text = ""
    for line in lines:
        add_para(tf, line, 13, MUTED, False, 8)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 12

    # 1 cover
    s = new_slide(prs)
    bar(s, Inches(0), Inches(0), W, H, DARK)
    bar(s, Inches(0), Inches(0), Inches(0.12), H, BLUE)
    add_text(s, Inches(0.7), Inches(1.9), Inches(11.5), Inches(0.4), "光途Work", 16, BLUE, True)
    add_text(s, Inches(0.7), Inches(2.35), Inches(11.8), Inches(1.1), "商业模式", 48, WHITE, True)
    add_text(
        s,
        Inches(0.7),
        Inches(3.55),
        Inches(11.2),
        Inches(0.9),
        "桌面客户端免费，控制面收费。卖本地 Agent 能力与网关额度，不卖软件授权，不卖上游 Key。",
        18,
        RGBColor(0xD6, 0xD3, 0xD1),
    )
    add_text(s, Inches(0.7), Inches(6.55), Inches(10), Inches(0.3), "2026  ·  内部讨论稿  ·  不含云端 Worker", 12, RGBColor(0xA8, 0xA2, 0x9E))

    # 2 positioning
    s = new_slide(prs)
    heading(s, "产品怎么赚钱", "一句话：用户买的是工作台使用权，请求走你们的网关。")
    items = [
        ("入口", "Windows 桌面客户端免费下载。Ask / Craft / Plan 都在本机工作空间完成。"),
        ("账本", "官网注册、买套餐、登录客户端。计量只发生在控制面，用户看不到厂商 Key。"),
        ("体验", "注册即领体验版，每种免费套餐只能领一次。用尽或到期后引导升级。"),
        ("可见", "购买后，客户端和官网只展示「已使用 xx%」，不展示 Token 数字。"),
    ]
    for i, (t, b) in enumerate(items):
        x = Inches(0.55) + (i % 2) * Inches(6.2)
        y = Inches(1.45) + (i // 2) * Inches(2.45)
        card_block(s, x, y, Inches(5.95), Inches(2.25), t, b, accent=(i == 3))
    footer(s, 2, total)

    # 3 sell / not sell
    s = new_slide(prs)
    heading(s, "卖什么，不卖什么", "商业边界决定技术边界。")
    sell = [
        "订阅：时间窗 + 后台额度上限",
        "能力开关：模型白名单、技能、MCP、专家团",
        "加油包：用尽后续量，不改档、不延期",
        "席位：团队版按人头扩容（落地后）",
        "视频生成：后续独立计量（见第 7 页）",
    ]
    nosell = [
        "上游 API Key 或 BYOK 个人通道",
        "云端 Worker / 关窗口继续跑任务",
        "向用户展示 Token 或单价账单",
        "技能市场分成（暂无供给端）",
        "把视频和对话额度 1:1 混用",
    ]
    rect(s, Inches(0.55), Inches(1.4), Inches(5.95), Inches(5.3), CARD, LINE)
    add_text(s, Inches(0.8), Inches(1.58), Inches(5.4), Inches(0.4), "卖", 18, BLUE, True)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(5.4), Inches(4.3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = ""
    for line in sell:
        add_para(tf, "·  " + line, 16, INK, False, 14)
    rect(s, Inches(6.75), Inches(1.4), Inches(5.95), Inches(5.3), CARD, LINE)
    add_text(s, Inches(7.0), Inches(1.58), Inches(5.4), Inches(0.4), "不卖", 18, WARN, True)
    box = s.shapes.add_textbox(Inches(7.0), Inches(2.1), Inches(5.4), Inches(4.3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = ""
    for line in nosell:
        add_para(tf, "·  " + line, 16, INK, False, 14)
    footer(s, 3, total)

    # 4 percentage UX
    s = new_slide(prs)
    heading(s, "用户只看见使用百分比", "Token 是进货账本，不是给用户看的商品名。")
    rows = [
        ("客户端工作台", "顶栏进度条 +「已使用 23%」。满 90% 变警示色。"),
        ("客户端账户页", "套餐名、已使用百分比、到期时间。可跳转官网升级。"),
        ("官网账户 / 订单", "同样只显示百分比。最近用量只留时间和模型名。"),
        ("套餐售卖页", "写「入门 / 标准 / 充裕额度」，不写「500 万 Token」。"),
        ("运营后台", "继续看 Token、成本、毛利。工作人员需要真实进货账。"),
        ("网关内部", "仍按 billed = 上游用量 × 倍率扣减，用尽返回 402。"),
    ]
    for i, (t, b) in enumerate(rows):
        y = Inches(1.35) + i * Inches(0.85)
        rect(s, Inches(0.55), y, Inches(12.2), Inches(0.75), CARD, LINE)
        add_text(s, Inches(0.8), y + Inches(0.18), Inches(2.6), Inches(0.42), t, 14, INK, True)
        add_text(s, Inches(3.5), y + Inches(0.18), Inches(8.9), Inches(0.42), b, 14, MUTED)
    footer(s, 4, total)

    # 5 packages
    s = new_slide(prs)
    heading(s, "三档套餐（对外口径）", "价格沿用现有种子数据。用户侧用额度档位，不用 Token 数字。")
    pkgs = [
        ("体验版", "免费 / 7 天", "入门额度\n仅 DeepSeek\n部分技能\n无 MCP / 专家团", "获客，验证工作流", False),
        ("专业版", "¥99 / 30 天", "标准额度\n多模型\n全技能 + MCP\n个人主力", "个人付费锚点", True),
        ("团队版", "¥299 / 30 天", "充裕额度\n全部已接入模型\nMCP + 专家团\n含 5 席（席位表待落地）", "小团队扩容", False),
    ]
    for i, (name, price, body, role, feat) in enumerate(pkgs):
        x = Inches(0.55) + i * Inches(4.15)
        r = rect(s, x, Inches(1.4), Inches(3.95), Inches(5.3), CARD, BLUE if feat else LINE)
        add_text(s, x + Inches(0.28), Inches(1.6), Inches(3.4), Inches(0.4), name + ("  · 推荐" if feat else ""), 18, INK, True)
        add_text(s, x + Inches(0.28), Inches(2.1), Inches(3.4), Inches(0.4), price, 22, BLUE if feat else INK, True)
        box = s.shapes.add_textbox(x + Inches(0.28), Inches(2.7), Inches(3.4), Inches(2.6))
        tf = box.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = ""
        for line in body.split("\n"):
            add_para(tf, line, 15, MUTED, False, 10)
        add_text(s, x + Inches(0.28), Inches(5.95), Inches(3.4), Inches(0.4), role, 13, BLUE, True)
    footer(s, 5, total)

    # 6 revenue mix
    s = new_slide(prs)
    heading(s, "四层收入（规划目标）", "稳态假设，不是历史账单。专业版应贡献一半以上经常性收入。")
    mix = [
        ("55%", "专业版订阅", "个人 ARR 主力。建议补年付 ¥990。"),
        ("25%", "团队版订阅", "更高额度 + 能力包。席位落地后才能按人头卖。"),
        ("15%", "额度加油包", "用尽后续量。建议 ≥ 套餐内含单价，保护订阅锚点。"),
        ("5%", "席位 / 视频", "加席 ¥49/人/月；视频单独配额，后续上线。"),
    ]
    for i, (pct, name, desc) in enumerate(mix):
        y = Inches(1.4) + i * Inches(1.25)
        rect(s, Inches(0.55), y, Inches(12.2), Inches(1.12), CARD, LINE)
        add_text(s, Inches(0.8), y + Inches(0.28), Inches(1.6), Inches(0.55), pct, 28, BLUE, True)
        add_text(s, Inches(2.7), y + Inches(0.22), Inches(3.2), Inches(0.35), name, 16, INK, True)
        add_text(s, Inches(2.7), y + Inches(0.58), Inches(9.6), Inches(0.35), desc, 14, MUTED)
    footer(s, 6, total)

    # 7 video
    s = new_slide(prs)
    heading(s, "后续：视频生成怎么赚钱", "视频进货远贵于文本，必须从第一天就和对话额度分开。")
    cards = [
        ("能力", "在专业版 / 团队版逐步开放视频模型。体验版默认不开，避免试用把成本打穿。"),
        ("计量", "独立「视频额度」：按成功生成条数或秒数扣减。用户侧仍只显示已使用百分比。"),
        ("售卖", "套餐内含少量视频条数 + 单独视频包。不要让用户用对话加油包去生成视频。"),
        ("风控", "上架时强制填写进货单价。倍率或条数成本必须覆盖最贵模型，运营可再加大。"),
    ]
    for i, (t, b) in enumerate(cards):
        x = Inches(0.55) + (i % 2) * Inches(6.2)
        y = Inches(1.4) + (i // 2) * Inches(2.5)
        card_block(s, x, y, Inches(5.95), Inches(2.3), t, b)
    footer(s, 7, total)

    # 8 funnel
    s = new_slide(prs)
    heading(s, "转化漏斗", "付费动机是继续完成手头任务，而不是抽象会员。")
    steps = [
        ("1", "注册", "邮箱验证码通过，自动发体验版，每种免费套餐一次。"),
        ("2", "试用", "7 天、弱模型、部分技能。日限额防止一夜刷光。"),
        ("3", "用尽", "网关 402。客户端与官网同时出现「已使用 100% / 去升级」。"),
        ("4", "付费", "浏览器打开支付宝 / 微信。不在 Electron 窗口内扣款。"),
        ("5", "续费", "到期提醒、80% 预警、加油包。把一次转化做成经常性收入。"),
    ]
    for i, (n, t, b) in enumerate(steps):
        y = Inches(1.35) + i * Inches(1.0)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), y + Inches(0.18), Inches(0.5), Inches(0.5))
        fill_shape(circ, BLUE)
        add_text(s, Inches(0.6), y + Inches(0.26), Inches(0.5), Inches(0.36), n, 14, WHITE, True, PP_ALIGN.CENTER)
        add_text(s, Inches(1.35), y + Inches(0.12), Inches(2.2), Inches(0.36), t, 16, INK, True)
        add_text(s, Inches(3.6), y + Inches(0.12), Inches(9.0), Inches(0.7), b, 14, MUTED)
    footer(s, 8, total)

    # 9 unit economics
    s = new_slide(prs)
    heading(s, "单位经济（内部核算）", "以下数字只出现在后台与本页，不出现在客户端或官网。")
    econ = [
        ("体验版", "收入 ¥0", "满额成本约 ¥0.2", "获客成本，可接受"),
        ("专业版 ¥99", "满额成本约 ¥10", "毛利率约 90%", "默认走 DeepSeek"),
        ("团队版 ¥299", "满额成本约 ¥40", "毛利率约 87%", "量折扣，靠席位放大"),
    ]
    for i, (a, b, c, d) in enumerate(econ):
        x = Inches(0.55) + i * Inches(4.15)
        card_block(s, x, Inches(1.4), Inches(3.95), Inches(2.5), a, f"{b}\n{c}\n{d}", accent=(i == 1))
    notes = [
        "后台扣费：billed = 上游用量 × 模型倍率。贵模型必须靠倍率烧得更快。",
        "现有代码用「2 分 / 千 billed」估成本，会低估真实毛利。上线前改为模型表里的真实进货价。",
        "唯一会打穿模式的事：把高价视频或旗舰文本模型勾进套餐，却不上调倍率 / 独立配额。",
    ]
    box = s.shapes.add_textbox(Inches(0.7), Inches(4.15), Inches(11.8), Inches(2.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = ""
    for n in notes:
        add_para(tf, "·  " + n, 15, INK, False, 12)
    footer(s, 9, total)

    # 10 tech
    s = new_slide(prs)
    heading(s, "技术实现分期", "钱只在 apps/api 确认。桌面端只读权益，官网负责下单。")
    phases = [
        ("P0  ·  1–2 周", "真实收款与漏损封口", "支付宝 / 微信异步回调；请求前预扣额度；并发闸；模型真实单价；用户 API 不返回 Token。"),
        ("P1  ·  2–4 周", "第二层收入", "加油包、年付、优惠码限次、升级继承剩余额度。账户页 80% 预警只显示百分比。"),
        ("P2  ·  按销售需要", "团队账本", "组织与成员表、共享额度池、发票状态机。席位未落地前，团队版按大额度个人套餐卖。"),
        ("P3  ·  视频上线时", "视频独立账本", "video_quota 或按条计数；套餐白名单；成功生成才扣；用户侧同样只给百分比。"),
    ]
    for i, (p, t, b) in enumerate(phases):
        y = Inches(1.35) + i * Inches(1.3)
        rect(s, Inches(0.55), y, Inches(12.2), Inches(1.18), CARD, LINE)
        add_text(s, Inches(0.8), y + Inches(0.18), Inches(3.3), Inches(0.32), p, 13, BLUE, True)
        add_text(s, Inches(4.2), y + Inches(0.16), Inches(8.2), Inches(0.32), t, 16, INK, True)
        add_text(s, Inches(4.2), y + Inches(0.55), Inches(8.2), Inches(0.48), b, 13, MUTED)
    footer(s, 10, total)

    # 11 not doing
    s = new_slide(prs)
    heading(s, "明确不做", "避免把产品做成云电脑或 Token 交易所。")
    nos = [
        ("云端 Worker", "任务只在用户本机执行。关窗口用托盘保持运行，不把工作空间同步到服务器。"),
        ("客户端贴 Key", "一旦 BYOK，计量账本被拆掉，毛利无法控。"),
        ("用户看见 Token", "对外只讲套餐与已使用百分比，避免比价和羊毛党按 Token 倒算。"),
        ("视频混用对话额度", "一条视频的进货可能吃掉大量对话额度，必须独立配额。"),
        ("未落地就卖 5 账号", "seats 字段还没有组织成员表。对外先别承诺多人共用。"),
        ("Electron 内支付", "收银台只在浏览器。桌面端只打开官网链接。"),
    ]
    for i, (t, b) in enumerate(nos):
        x = Inches(0.55) + (i % 2) * Inches(6.2)
        y = Inches(1.35) + (i // 2) * Inches(1.75)
        card_block(s, x, y, Inches(5.95), Inches(1.6), t, b)
    footer(s, 11, total)

    # 12 decisions
    s = new_slide(prs)
    heading(s, "建议立刻拍板的决策", "可以改价格，不建议改结构。")
    decisions = [
        ("计费对象", "后台按额度扣；用户只看已使用百分比。"),
        ("默认模型", "DeepSeek，保证试用和专业版主力成本。"),
        ("支付", "官网收银台；桌面端只给链接。"),
        ("执行位置", "仅本机。云端 Worker 已从产品移除。"),
        ("视频", "后续独立配额，不与对话额度 1:1 混用。"),
        ("团队版", "席位表上线前，按大额度个人套餐销售。"),
        ("数据库", "先 SQLite 跑通支付；付费用户过千再迁 Postgres。"),
        ("体验版", "继续自动发放，且每人每种免费套餐一次。"),
    ]
    for i, (t, b) in enumerate(decisions):
        col = i % 2
        row = i // 2
        x = Inches(0.55) + col * Inches(6.2)
        y = Inches(1.32) + row * Inches(1.28)
        rect(s, x, y, Inches(5.95), Inches(1.15), CARD, LINE)
        add_text(s, x + Inches(0.22), y + Inches(0.16), Inches(5.5), Inches(0.32), t, 14, BLUE, True)
        add_text(s, x + Inches(0.22), y + Inches(0.52), Inches(5.5), Inches(0.48), b, 13, MUTED)
    footer(s, 12, total)

    out = Path(__file__).resolve().parent / "光途Work商业模式.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
